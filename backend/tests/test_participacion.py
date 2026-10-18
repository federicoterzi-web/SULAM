"""División Participación y Sensibilización: tipos, eventos en la agenda, campañas, asistentes, fotos, brief y estadísticas."""

from __future__ import annotations

import io
from datetime import datetime, timedelta, timezone

from conftest import AREA_ADMISION, crear_usuario

DIV = "División Participación y Sensibilización Ciudadana"
ROL = "Participación y Sensibilización"


def _foto(color=(200, 80, 60)) -> bytes:
    from PIL import Image

    b = io.BytesIO()
    Image.new("RGB", (640, 480), color).save(b, "JPEG")
    return b.getvalue()


def _cuando(dias=1, horas=0):
    return (datetime.now(timezone.utc) + timedelta(days=dias, hours=horas)).replace(microsecond=0).isoformat()


def _tipo(c, nombre="Charla"):
    r = c.post("/api/participacion/tipos", json={"nombre": nombre})
    assert r.status_code == 201, r.text
    return r.json()["id"]


def _evento(c, tipo_id, **kw):
    datos = {"tipo_id": tipo_id, "inicio": _cuando(), "lugar": "Centro de jubilados"}
    datos.update(kw)
    r = c.post("/api/participacion/eventos", json=datos)
    assert r.status_code == 201, r.text
    return r.json()


def _realizar(c, evento_id):
    r = c.post(f"/api/agenda/{evento_id}/resultado", json={"estado": "realizado", "resultado": "Muy buena convocatoria."})
    assert r.status_code == 200, r.text


def test_permisos_del_area(admin, cliente):
    crear_usuario("part", ROL, DIV)
    crear_usuario("tec", "Técnico", AREA_ADMISION)
    crear_usuario("dir", "Directora", "Dirección de Adultos Mayores")
    p, t, d = cliente("part"), cliente("tec"), cliente("dir")
    # El área define sus tipos; quien no pertenece al área no.
    tid = _tipo(p, "Charla")
    assert t.post("/api/participacion/tipos", json={"nombre": "Otra"}).status_code == 403
    assert t.get("/api/participacion/tipos").status_code == 403
    assert p.post("/api/participacion/tipos", json={"nombre": "charla"}).status_code == 409  # mismo nombre
    assert d.post("/api/participacion/tipos", json={"nombre": "Otra"}).status_code == 403  # la directora solo consulta
    assert d.get("/api/participacion/tipos").status_code == 200
    # Un tipo dado de baja no se puede elegir para eventos nuevos.
    assert p.patch(f"/api/participacion/tipos/{tid}", json={"activo": False}).status_code == 200
    assert p.post("/api/participacion/eventos", json={"tipo_id": tid, "inicio": _cuando()}).status_code == 400
    # El rol del área también arma el mapa de recursos.
    assert "recursos.crear" in p.yo["permisos"] or any("recursos" in k for k in p.yo["permisos"])


def test_evento_titulo_agenda_y_resumen_semanal(admin, cliente):
    crear_usuario("part", ROL, DIV)
    crear_usuario("dir", "Directora", "Dirección de Adultos Mayores")
    p, d = cliente("part"), cliente("dir")
    charla = _tipo(p, "Charla")
    ev = _evento(p, charla)
    assert ev["titulo"] == "Charla" and ev["tipo"] == "evento"  # sin flyer: el título es el tipo de evento
    con_flyer = _evento(
        p, charla, nombre_flyer="Cuidar sin cansarse", instituciones=["Centro de Jubilados 9 de Julio", "CAPS Sur", "caps sur"],
        resumen="Charla abierta sobre cuidados.", sugerir_directora=True, organiza_otra="Ministerio de Salud", inicio=_cuando(2),
    )
    assert con_flyer["titulo"] == "Cuidar sin cansarse"
    assert con_flyer["participacion"]["instituciones"] == ["Centro de Jubilados 9 de Julio", "CAPS Sur"]
    # Aparece en la agenda común, pero no se crea ni se edita por la vía general.
    desde, hasta = _cuando(-1), _cuando(10)
    en_agenda = p.get("/api/agenda", params={"desde": desde, "hasta": hasta}).json()["eventos"]
    assert {e["titulo"] for e in en_agenda} == {"Charla", "Cuidar sin cansarse"}
    assert p.post("/api/agenda", json={"tipo": "evento", "titulo": "Xyz", "inicio": _cuando()}).status_code == 400
    assert p.patch(f"/api/agenda/{con_flyer['id']}", json={"titulo": "Otro"}).status_code == 400
    assert "evento" not in [t["clave"] for t in p.get("/api/agenda/tipos").json()]
    # Resumen semanal para la Dirección.
    lunes = (datetime.now() + timedelta(days=1)).date()
    s = d.get("/api/participacion/semana", params={"lunes": lunes.isoformat()}).json()
    assert s["total"] == 2 and s["con_presencia_sugerida"] == 1
    assert "Cuidar sin cansarse" in s["texto"] and "Se sugiere la presencia de la Dirección" in s["texto"] and "CAPS Sur" in s["texto"]
    assert sum(len(x["eventos"]) for x in s["dias"]) == 2
    # El tablero de la Dirección lo trae y el del área también.
    assert d.get("/api/inicio").json()["datos"]["participacion"]["total"] >= 0
    assert p.get("/api/inicio").json()["perfil"] == "participacion"
    assert "semana" in p.get("/api/inicio").json()["datos"]


def test_edicion_en_serie_y_cambio_de_nombre(admin, cliente):
    crear_usuario("part", ROL, DIV)
    p = cliente("part")
    t = _tipo(p, "Taller")
    r = _evento(p, t, repetir={"frecuencia": "semanal", "cantidad": 3}, instituciones=["CAPS Norte"])
    assert r["creadas"] == 3
    desde, hasta = _cuando(-1), _cuando(40)
    filas = p.get("/api/agenda", params={"desde": desde, "hasta": hasta}).json()["eventos"]
    assert len(filas) == 3
    m = p.patch(f"/api/participacion/eventos/{filas[1]['id']}", json={"nombre_flyer": "Memoria activa", "instituciones": ["Hogar San José"], "aplicar_a": "siguientes", "motivo": "Se definió el flyer"})
    assert m.status_code == 200, m.text
    filas = p.get("/api/agenda", params={"desde": desde, "hasta": hasta}).json()["eventos"]
    assert [f["titulo"] for f in filas] == ["Taller", "Memoria activa", "Memoria activa"]
    assert p.get(f"/api/participacion/eventos/{filas[2]['id']}").json()["participacion"]["instituciones"] == ["Hogar San José"]
    assert p.get(f"/api/participacion/eventos/{filas[0]['id']}").json()["participacion"]["instituciones"] == ["CAPS Norte"]
    h = p.get(f"/api/agenda/{filas[2]['id']}").json()["historial"]
    assert h and h[0]["motivo"] == "Se definió el flyer"


def test_realizado_asistentes_fotos_y_brief(admin, cliente):
    from docx import Document

    crear_usuario("part", ROL, DIV)
    p = cliente("part")
    t = _tipo(p, "Feria")
    ev = _evento(p, t, nombre_flyer="Feria de la Salud", resumen="Chequeos y charlas.", instituciones=["PAMI"])
    # Antes de realizarse no hay asistentes, fotos ni brief.
    assert p.patch(f"/api/participacion/eventos/{ev['id']}/asistentes", json={"asistentes": 50}).status_code == 409
    assert p.post(f"/api/participacion/eventos/{ev['id']}/fotos", files={"archivo": ("a.jpg", _foto(), "image/jpeg")}).status_code == 409
    assert p.get(f"/api/participacion/eventos/{ev['id']}/brief.docx").status_code == 409
    _realizar(p, ev["id"])
    r = p.patch(f"/api/participacion/eventos/{ev['id']}/asistentes", json={"asistentes": 85})
    assert r.status_code == 200 and r.json()["participacion"]["asistentes"] == 85
    assert p.patch(f"/api/participacion/eventos/{ev['id']}/asistentes", json={"asistentes": -3}).status_code == 422
    f1 = p.post(f"/api/participacion/eventos/{ev['id']}/fotos", files={"archivo": ("a.jpg", _foto(), "image/jpeg")})
    f2 = p.post(f"/api/participacion/eventos/{ev['id']}/fotos", files={"archivo": ("b.jpg", _foto((20, 90, 160)), "image/jpeg")})
    assert f1.status_code == 201 and f2.status_code == 201 and {f1.json()["orden"], f2.json()["orden"]} == {1, 2}
    # Dos fotos como máximo: para subir otra hay que quitar una.
    assert p.post(f"/api/participacion/eventos/{ev['id']}/fotos", files={"archivo": ("c.jpg", _foto(), "image/jpeg")}).status_code == 409
    assert p.post(f"/api/participacion/eventos/{ev['id']}/fotos", files={"archivo": ("x.txt", b"hola", "text/plain")}).status_code in (409, 415)
    assert p.get(f"/api/participacion/eventos/{ev['id']}/fotos/{f1.json()['id']}").content == _foto()
    assert p.delete(f"/api/participacion/eventos/{ev['id']}/fotos/{f2.json()['id']}").status_code == 204
    assert p.post(f"/api/participacion/eventos/{ev['id']}/fotos", files={"archivo": ("c.png", _png(), "image/png")}).status_code == 201
    # Brief en Word con los datos y las dos fotos.
    b = p.get(f"/api/participacion/eventos/{ev['id']}/brief.docx")
    assert b.status_code == 200 and b.headers["content-type"].endswith("wordprocessingml.document")
    doc = Document(io.BytesIO(b.content))
    texto = "\n".join(x.text for x in doc.paragraphs) + "\n".join(c.text for t_ in doc.tables for row in t_.rows for c in row.cells)
    assert "Feria de la Salud" in texto and "85" in texto and "PAMI" in texto and "Chequeos y charlas." in texto and "Muy buena convocatoria." in texto
    assert len(doc.inline_shapes) == 2


def _png() -> bytes:
    from PIL import Image

    b = io.BytesIO()
    Image.new("RGB", (300, 200), (10, 120, 60)).save(b, "PNG")
    return b.getvalue()


def test_campania_con_material_eventos_y_brief(admin, cliente):
    from docx import Document

    crear_usuario("part", ROL, DIV)
    crear_usuario("tec", "Técnico", AREA_ADMISION)
    p, t = cliente("part"), cliente("tec")
    tipo = _tipo(p, "Jornada")
    c = p.post("/api/participacion/campanias", json={"nombre": "Buen trato", "descripcion": "Campaña de sensibilización.", "desde": "2030-05-01", "hasta": "2030-06-30"})
    assert c.status_code == 201, c.text
    cid = c.json()["id"]
    # Material: PDF, imagen y Word sí; un ejecutable disfrazado no.
    assert p.post(f"/api/participacion/campanias/{cid}/materiales", files={"archivo": ("flyer.png", _png(), "image/png")}).status_code == 201
    pdf = p.post(f"/api/participacion/campanias/{cid}/materiales", files={"archivo": ("guia.pdf", b"%PDF-1.4\n%%EOF", "application/pdf")})
    assert pdf.status_code == 201
    assert p.post(f"/api/participacion/campanias/{cid}/materiales", files={"archivo": ("virus.exe", b"MZ....", "application/octet-stream")}).status_code == 415
    assert p.post(f"/api/participacion/campanias/{cid}/materiales", files={"archivo": ("falso.pdf", b"no es pdf", "application/pdf")}).status_code == 415
    assert p.get(f"/api/participacion/campanias/{cid}/materiales/{pdf.json()['id']}").content.startswith(b"%PDF")
    # Eventos de la campaña: uno propio y otro a cargo de otra organización con presencia sugerida.
    e1 = _evento(p, tipo, nombre_flyer="Apertura", campania_id=cid, inicio=_cuando(-3))
    e2 = _evento(p, tipo, nombre_flyer="Cierre", campania_id=cid, organiza_otra="Municipalidad de Yerba Buena", sugerir_directora=True, inicio=_cuando(20))
    _realizar(p, e1["id"])
    p.patch(f"/api/participacion/eventos/{e1['id']}/asistentes", json={"asistentes": 120})
    det = p.get(f"/api/participacion/campanias/{cid}").json()
    assert det["eventos"] == 2 and det["realizados"] == 1 and det["asistentes"] == 120 and len(det["lista_materiales"]) == 2
    assert [e["organiza_otra"] for e in det["lista_eventos"]] == [None, "Municipalidad de Yerba Buena"]
    brief = p.get(f"/api/participacion/campanias/{cid}/brief.docx")
    assert brief.status_code == 200
    doc = Document(io.BytesIO(brief.content))
    texto = "\n".join(x.text for x in doc.paragraphs)
    assert "Buen trato" in texto and "Apertura" in texto and "Cierre" in texto and "Municipalidad de Yerba Buena" in texto
    # Quien no es del área no ve la campaña ni su material.
    assert t.get(f"/api/participacion/campanias/{cid}").status_code == 403
    assert p.patch(f"/api/participacion/campanias/{cid}", json={"estado": "cerrada"}).json()["estado"] == "cerrada"
    assert p.delete(f"/api/participacion/campanias/{cid}/materiales/{pdf.json()['id']}").status_code == 204


def test_estadisticas_de_participacion(admin, cliente):
    crear_usuario("part", ROL, DIV)
    crear_usuario("tec", "Técnico", AREA_ADMISION)
    p, t = cliente("part"), cliente("tec")
    charla, taller = _tipo(p, "Charla"), _tipo(p, "Taller")
    c = p.post("/api/participacion/campanias", json={"nombre": "Buen trato"}).json()["id"]
    a = _evento(p, charla, campania_id=c, inicio=_cuando(-5), instituciones=["PAMI", "CAPS Sur"])
    b = _evento(p, taller, inicio=_cuando(-4), organiza_otra="ONG Vida", instituciones=["PAMI"])
    _evento(p, taller, inicio=_cuando(-3))  # pasó y no tiene resultado
    for e, n in ((a, 40), (b, 60)):
        _realizar(p, e["id"])
        p.patch(f"/api/participacion/eventos/{e['id']}/asistentes", json={"asistentes": n})
    hoy = datetime.now().date()
    params = {"desde": (hoy - timedelta(days=30)).isoformat(), "hasta": (hoy + timedelta(days=5)).isoformat()}
    d = p.get("/api/estadisticas/participacion", params=params).json()
    r = {x["clave"]: x["valor"] for x in d["resumen"]}
    assert r["eventos"] == 3 and r["realizados"] == 2 and r["asistentes"] == 100 and r["promedio"] == 50.0 and r["otras"] == 1 and r["sin_resultado"] == 1
    tablas = {x["clave"]: x["filas"] for x in d["tablas"]}
    assert {f["etiqueta"]: f["asistentes"] for f in tablas["por_tipo"]} == {"Charla": 40, "Taller": 60}
    assert {f["etiqueta"]: f["cantidad"] for f in tablas["instituciones"]}["PAMI"] == 2
    assert p.get("/api/estadisticas/participacion.xlsx", params=params).status_code == 200
    pdf = p.get("/api/estadisticas/participacion.pdf", params=params)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    # Sin permiso de Participación ni de Estadísticas: no.
    assert t.get("/api/estadisticas/participacion", params=params).status_code == 403
    # La Dirección lo ve todo.
    assert admin.get("/api/estadisticas/participacion", params=params).json()["resumen"][0]["valor"] == 3


def test_roles_existentes_reciben_permisos_de_participacion(db):
    from sqlalchemy import delete, select

    from sulam.models import Ajuste, Rol, RolPermiso
    from sulam.seed import inicializar_catalogos

    dire = db.execute(select(Rol).where(Rol.nombre == "Directora")).scalar_one()
    db.execute(delete(RolPermiso).where(RolPermiso.rol_id == dire.id, RolPermiso.modulo == "participacion"))
    db.execute(delete(Ajuste).where(Ajuste.clave == "permisos_iniciales_participacion"))
    db.commit()
    db.expire_all()
    inicializar_catalogos(db)
    db.commit()
    dire = db.execute(select(Rol).where(Rol.nombre == "Directora")).scalar_one()
    assert {(x.accion, x.alcance) for x in dire.permisos if x.modulo == "participacion"} == {("consultar", "todos"), ("descargar", "todos")}


def test_brief_en_pdf_de_evento_y_campania(admin, cliente):
    crear_usuario("part", ROL, DIV)
    p = cliente("part")
    ev = _evento(p, _tipo(p, "Feria"), nombre_flyer="Feria de la Salud", resumen="Chequeos.")
    camp = p.post("/api/participacion/campanias", json={"nombre": "Octubre Rosa"}).json()
    p.patch(f"/api/participacion/eventos/{ev['id']}", json={"campania_id": camp["id"]})
    assert p.get(f"/api/participacion/eventos/{ev['id']}/brief.pdf").status_code == 409  # todavía no se realizó
    _realizar(p, ev["id"])
    p.post(f"/api/participacion/eventos/{ev['id']}/fotos", files={"archivo": ("a.jpg", _foto(), "image/jpeg")})
    p.post(f"/api/participacion/eventos/{ev['id']}/fotos", files={"archivo": ("b.png", _png(), "image/png")})
    for url in (f"/api/participacion/eventos/{ev['id']}/brief.pdf", f"/api/participacion/campanias/{camp['id']}/brief.pdf"):
        r = p.get(url)
        assert r.status_code == 200 and r.headers["content-type"] == "application/pdf" and r.content.startswith(b"%PDF")
        assert ".pdf" in r.headers["content-disposition"]
    assert admin.get(f"/api/participacion/eventos/{ev['id']}/brief.pdf").status_code == 200


def test_aviso_de_reprogramacion(admin, cliente):
    crear_usuario("part", ROL, DIV)
    crear_usuario("part2", ROL, DIV)
    crear_usuario("dir", "Directora", "Dirección de Adultos Mayores")
    p, otro, d = cliente("part"), cliente("part2"), cliente("dir")
    t = _tipo(p, "Charla")
    ev = _evento(p, t, sugerir_directora=True)
    sin_aviso = _evento(p, t)
    avisos = lambda c: c.get("/api/avisos").json()["avisos"].get("participacion_cambios", [])
    assert avisos(p) == [] and avisos(d) == []
    # La reprogramación la hace otra persona del área: avisa a quien lo creó y a la Directora (se sugiere su presencia).
    r = otro.patch(f"/api/participacion/eventos/{ev['id']}", json={"inicio": _cuando(3), "aplicar_a": "esta"})
    assert r.status_code == 200, r.text
    a1 = avisos(p)
    assert len(a1) == 1 and a1[0].startswith(ev["id"]) and len(avisos(d)) == 1
    assert avisos(otro) == []  # no se avisa de los cambios propios
    # Un evento sin presencia sugerida no le llega a la Directora.
    otro.patch(f"/api/participacion/eventos/{sin_aviso['id']}", json={"lugar": "Otro salón", "aplicar_a": "esta"})
    assert len(avisos(p)) == 2 and len(avisos(d)) == 1
    # Un cambio nuevo del mismo evento vuelve a avisar (cambia el identificador).
    otro.patch(f"/api/participacion/eventos/{ev['id']}", json={"lugar": "Plaza", "aplicar_a": "esta"})
    assert set(avisos(p)) != set(a1) and len(avisos(d)) == 1


def test_teselas_del_mapa_pasan_por_el_servidor(admin, monkeypatch):
    import urllib.request

    pedidos = []

    class Resp:
        headers = type("H", (), {"get_content_type": lambda s: "image/png"})()
        def __enter__(self): return self
        def __exit__(self, *a): return False
        def read(self, n=-1): return b"\x89PNG-demo"

    def falso(pedido, timeout=0):
        pedidos.append(pedido)
        return Resp()

    monkeypatch.setattr("urllib.request.urlopen", falso)
    assert admin.get("/api/mapa/config").json()["teselas"] == "/api/mapa/teselas/{z}/{x}/{y}.png"
    r = admin.get("/api/mapa/teselas/12/1300/2500.png")
    assert r.status_code == 200 and r.content == b"\x89PNG-demo"
    assert pedidos[0].get_header("User-agent").startswith("SULAM/")
    admin.get("/api/mapa/teselas/12/1300/2500.png")
    assert len(pedidos) == 1  # la segunda sale del disco
    assert admin.get("/api/mapa/teselas/40/1/1.png").status_code == 404
