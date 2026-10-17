"""Estadísticas de casos y de funcionamiento interno."""

from __future__ import annotations

import json
from datetime import date

from conftest import AREA_ADMISION, AREA_PROMOCION, crear_usuario, nueva_persona, nuevo_caso


def _preparar(admin):
    areas = {a["nombre"]: a["id"] for a in admin.get("/api/catalogos").json()["areas"]}
    crear_usuario("tec", "Técnico", AREA_ADMISION)
    crear_usuario("jefa", "Jefe/a de departamento", AREA_ADMISION)
    tec = {u["nombre_completo"]: u["id"] for u in admin.get("/api/catalogos").json()["usuarios_asignables"]}["Tec"]
    p1 = nueva_persona(admin, apellidos="UNO", numero_documento="1111111", fecha_nacimiento="1940-05-01", genero="F", departamento="Capital")
    p2 = nueva_persona(admin, apellidos="DOS", nombres="Otro", numero_documento="2222222", fecha_nacimiento="1958-01-01", genero="M")
    p3 = nueva_persona(admin, apellidos="TRES", nombres="Otra", numero_documento="3333333")
    c1 = nuevo_caso(admin, p1["id"], area_id=areas[AREA_ADMISION], responsable_id=tec, indicadores=["maltrato.fisico", "negligencia"])
    c2 = nuevo_caso(admin, p2["id"], area_id=areas[AREA_ADMISION], indicadores=["vulnerabilidad.habitacional"])
    c3 = nuevo_caso(admin, p3["id"], area_id=areas[AREA_PROMOCION])
    admin.post(f"/api/casos/{c2['id']}/estado", json={"estado": "cerrado", "motivo": "Intervención finalizada"})
    admin.post(f"/api/casos/{c1['id']}/evolucion", json={"texto": "Visita domiciliaria (prueba)"})
    return areas, tec, (c1, c2, c3)


def _tabla(d, clave):
    for s in d["secciones"]:
        for t in s["tablas"]:
            if t["clave"] == clave:
                return {f["etiqueta"]: f for f in t["filas"]}
    raise KeyError(clave)


def test_estadisticas_de_casos(admin):
    areas, tec, (c1, c2, c3) = _preparar(admin)
    d = admin.get("/api/estadisticas/casos").json()
    res = {r["clave"]: r["valor"] for r in d["resumen"]}
    assert res["ingresados"] == 3 and res["maltrato"] == 1 and res["cerrados"] == 1 and res["activos"] == 2
    ind = _tabla(d, "indicadores")
    assert ind["Maltrato"]["cantidad"] == 1 and ind["Negligencia"]["cantidad"] == 1 and ind["Sin dato"]["cantidad"] == 1
    assert _tabla(d, "indicadores_detalle")["Maltrato: Físico"]["cantidad"] == 1
    assert _tabla(d, "edad")["85 a 89"]["cantidad"] == 1 and _tabla(d, "edad")["Sin dato"]["cantidad"] == 1
    assert _tabla(d, "responsable")["Tec"]["cantidad"] == 1
    assert sum(m["ingresados"] for m in d["mensual"]) == 3

    # El filtro de cada fila reproduce el número en la lista de casos.
    for clave in ("indicadores", "edad", "area", "responsable", "estado", "genero", "departamento"):
        for fila in _tabla(d, clave).values():
            if not fila.get("filtro"):
                continue
            f = json.dumps(d["base"] + [d["periodo_filtro"]] + fila["filtro"])
            r = admin.get("/api/casos", params={"f": f, "con_total": True}).json()
            assert r["total"] == fila["cantidad"], (clave, fila)

    # Con filtros y por área
    f = json.dumps([{"c": "caso.indicadores", "o": "alguno", "v": ["maltrato"]}])
    assert admin.get("/api/estadisticas/casos", params={"f": f}).json()["total"] == 1
    assert admin.get("/api/estadisticas/casos", params={"area_id": areas[AREA_PROMOCION]}).json()["total"] == 1


def test_alcance_de_jefaturas_y_personal(admin, cliente):
    areas, _, _ = _preparar(admin)
    jefa = cliente("jefa")
    d = jefa.get("/api/estadisticas/casos").json()
    assert d["total"] == 2  # solo su área
    assert jefa.get("/api/estadisticas/casos", params={"area_id": areas[AREA_PROMOCION]}).status_code == 403
    assert jefa.get("/api/estadisticas/interno").status_code == 200
    crear_usuario("tecdos", "Técnico", AREA_ADMISION)
    assert cliente("tecdos").get("/api/estadisticas/casos").status_code == 403
    crear_usuario("pers", "Personal")
    p = cliente("pers")
    assert p.get("/api/estadisticas/casos").status_code == 403
    i = p.get("/api/estadisticas/interno").json()
    assert i["asistencia"] is not None and i["carga"] is None and i["circuitos"] is None


def test_funcionamiento_interno(admin):
    _preparar(admin)
    hoy = date.today().isoformat()
    admin.post("/api/novedades", data={"titulo": "Circular de prueba", "cuerpo": "<p>Texto</p>", "requiere_confirmacion": "true"})
    d = admin.get("/api/estadisticas/interno", params={"desde": hoy[:8] + "01", "hasta": hoy}).json()
    carga = {r["clave"]: r["valor"] for r in d["carga"]["resumen"]}
    assert carga["activos"] == 2 and carga["sin_responsable"] == 1
    tecnicos = {f["etiqueta"]: f for f in d["carga"]["tablas"][0]["filas"]}
    assert tecnicos["Tec"]["activos"] == 1
    assert d["asistencia"] is not None and d["circuitos"]["resumen"]
    assert {r["clave"] for r in d["comunicacion"]["resumen"]} >= {"novedades", "lectura"}


def test_descargas_excel_y_pdf(admin, cliente):
    import io

    from openpyxl import load_workbook

    _preparar(admin)
    f = json.dumps([{"c": "persona.genero", "o": "es", "v": ["F"]}])
    r = admin.get("/api/estadisticas/casos.xlsx", params={"f": f})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Resumen" and "Detalle de casos" in wb.sheetnames
    assert wb["Detalle de casos"].max_row == 2  # una sola persona de género F
    assert "Género es Femenino" in str(wb["Resumen"]["B4"].value)
    r = admin.get("/api/estadisticas/interno.xlsx")
    assert r.status_code == 200 and "Carga de trabajo y tiempos" in load_workbook(io.BytesIO(r.content)).sheetnames
    for s in ("casos", "interno"):
        r = admin.get(f"/api/estadisticas/{s}.pdf")
        assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert admin.get("/api/estadisticas/otra.pdf").status_code == 404
    crear_usuario("tecx", "Técnico", AREA_ADMISION)
    assert cliente("tecx").get("/api/estadisticas/casos.xlsx").status_code == 403
    hoy = date.today().isoformat()
    r = admin.get("/api/personal/incidencias.xlsx", params={"desde": hoy[:8] + "01", "hasta": hoy})
    assert r.status_code == 200
    assert [c.value for c in load_workbook(io.BytesIO(r.content))["Incidencias"][1]][:3] == ["Trabajador", "Fecha", "Turno desde"]
