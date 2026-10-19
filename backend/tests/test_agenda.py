"""Agenda: actividades propias y del área, resultado y su paso a la evolución del caso."""

from __future__ import annotations

from datetime import datetime, timedelta, timezone

from conftest import AREA_ADMISION, crear_usuario, nueva_persona, nuevo_caso

OTRA = "Departamento Promoción Integral de Derechos"


def _cuando(horas=24, dias=0):
    return (datetime.now(timezone.utc) + timedelta(hours=horas, days=dias)).replace(microsecond=0).isoformat()


def _ev(c, **kw):
    datos = {"tipo": "visita", "titulo": "Visita domiciliaria", "inicio": _cuando(), "lugar": "Domicilio"}
    datos.update(kw)
    return c.post("/api/agenda", json=datos)


def test_propia_y_del_area(admin, cliente):
    crear_usuario("tec1", "Técnico", AREA_ADMISION)
    crear_usuario("tec2", "Técnico", AREA_ADMISION)
    crear_usuario("otro", "Técnico", OTRA)
    a, b, o = cliente("tec1"), cliente("tec2"), cliente("otro")
    e = _ev(a, titulo="Reunión de equipo", tipo="reunion")
    assert e.status_code == 201, e.text
    desde, hasta = _cuando(-48), _cuando(24 * 10)
    # La ve quien participa; el compañero de área la ve en «el área» pero no en «las mías»; otra área no la ve.
    assert len(a.get("/api/agenda", params={"desde": desde, "hasta": hasta}).json()["eventos"]) == 1
    assert b.get("/api/agenda", params={"desde": desde, "hasta": hasta}).json()["eventos"] == []
    assert len(b.get("/api/agenda", params={"desde": desde, "hasta": hasta, "ambito": "area"}).json()["eventos"]) == 1
    assert o.get("/api/agenda", params={"desde": desde, "hasta": hasta, "ambito": "area"}).json()["eventos"] == []
    assert o.get(f"/api/agenda/{e.json()['id']}").status_code == 403
    # Con alcance «propios» para modificar, el compañero puede verla pero no tocarla.
    assert b.patch(f"/api/agenda/{e.json()['id']}", json={"titulo": "Cambio ajeno"}).status_code == 403
    assert a.patch(f"/api/agenda/{e.json()['id']}", json={"titulo": "Reunión de equipo (aula)"}).json()["titulo"].endswith("(aula)")


def test_no_agenda_para_otra_area_ni_a_personas_dadas_de_baja(admin, cliente):
    crear_usuario("tec1", "Técnico", AREA_ADMISION)
    crear_usuario("otro", "Técnico", OTRA)
    a = cliente("tec1")
    otro_id = next(u["id"] for u in admin.get("/api/admin/usuarios").json() if u["usuario"] == "otro")
    r = _ev(a, participantes=[otro_id])
    assert r.status_code == 403, r.text
    assert _ev(a, tipo="inexistente").status_code == 422
    assert _ev(a, fin=_cuando(-5)).status_code == 422


def test_resultado_de_una_salida_queda_en_el_caso(admin, cliente):
    crear_usuario("tec1", "Técnico", AREA_ADMISION)
    a = cliente("tec1")
    p = nueva_persona(admin)
    caso = nuevo_caso(admin, p["id"])
    e = _ev(admin, tipo="salida", titulo="Salida a domicilio", persona_id=p["id"], caso_id=caso["id"], lugar="Barrio Sur").json()
    assert e["persona"]["nombre"].startswith("PRUEBA") and e["caso"]["id"] == caso["id"]
    # No se puede dar por no realizada sin motivo.
    assert admin.post(f"/api/agenda/{e['id']}/resultado", json={"estado": "no_realizado"}).status_code == 400
    r = admin.post(f"/api/agenda/{e['id']}/resultado", json={"estado": "realizado", "resultado": "Se la encontró bien; se acordó nueva visita."})
    assert r.status_code == 200 and r.json()["estado"] == "realizado"
    ficha = admin.get(f"/api/casos/{caso['id']}").json()
    textos = [x["texto"] for x in ficha.get("evoluciones", [])]
    assert any("Salida realizada" in t and "nueva visita" in t for t in textos), ficha.keys()
    # Ya tiene resultado: no se vuelve a resolver ni se modifica.
    assert admin.post(f"/api/agenda/{e['id']}/resultado", json={"estado": "cancelado"}).status_code == 409
    assert admin.patch(f"/api/agenda/{e['id']}", json={"titulo": "otro"}).status_code == 409


def test_avisos_de_agenda_solo_para_participantes(admin, cliente):
    crear_usuario("tec1", "Técnico", AREA_ADMISION)
    crear_usuario("tec2", "Técnico", AREA_ADMISION)
    a, b = cliente("tec1"), cliente("tec2")
    assert _ev(a, inicio=_cuando(3)).status_code == 201
    assert _ev(a, titulo="Lejana", inicio=_cuando(24 * 5)).status_code == 201
    assert len(a.get("/api/avisos").json()["avisos"]["agenda"]) == 1
    assert b.get("/api/avisos").json()["avisos"]["agenda"] == []


def test_roles_existentes_reciben_los_permisos_de_agenda_una_sola_vez(db):
    from sqlalchemy import delete, select

    from sulam.models import Ajuste, Rol, RolPermiso
    from sulam.seed import inicializar_catalogos

    tec = db.execute(select(Rol).where(Rol.nombre == "Técnico")).scalar_one()
    db.execute(delete(RolPermiso).where(RolPermiso.rol_id == tec.id, RolPermiso.modulo == "agenda"))
    db.execute(delete(Ajuste).where(Ajuste.clave == "permisos_iniciales_agenda"))
    db.commit()
    db.expire_all()
    inicializar_catalogos(db)
    db.commit()
    tec = db.execute(select(Rol).where(Rol.nombre == "Técnico")).scalar_one()
    assert {(p.accion, p.alcance) for p in tec.permisos if p.modulo == "agenda"} == {("consultar", "area"), ("crear", "area"), ("modificar", "propios")}
    # Si el administrador los quita, no se vuelven a poner.
    db.execute(delete(RolPermiso).where(RolPermiso.rol_id == tec.id, RolPermiso.modulo == "agenda"))
    db.commit()
    db.expire_all()
    inicializar_catalogos(db)
    db.commit()
    db.expire_all()
    assert not [p for p in db.execute(select(Rol).where(Rol.nombre == "Técnico")).scalar_one().permisos if p.modulo == "agenda"]


def test_actividad_repetida_se_edita_y_se_cancela_desde_una_en_adelante(admin):
    r = _ev(admin, tipo="reunion", titulo="Reunión de equipo", inicio=_cuando(24), repetir={"frecuencia": "semanal", "cantidad": 4})
    assert r.status_code == 201, r.text
    assert r.json()["creadas"] == 4 and r.json()["repeticion"] == "semanal"
    desde, hasta = _cuando(-24), _cuando(24 * 40)
    filas = admin.get("/api/agenda", params={"desde": desde, "hasta": hasta}).json()["eventos"]
    assert len(filas) == 4 and len({f["serie_id"] for f in filas}) == 1
    segunda = filas[1]
    # Se pasa de hora desde la segunda en adelante: la primera queda como estaba.
    nuevo = (datetime.fromisoformat(segunda["inicio"]) + timedelta(hours=2)).isoformat()
    m = admin.patch(f"/api/agenda/{segunda['id']}", json={"inicio": nuevo, "aplicar_a": "siguientes", "motivo": "Cambió la sala"})
    assert m.status_code == 200, m.text
    filas = admin.get("/api/agenda", params={"desde": desde, "hasta": hasta}).json()["eventos"]
    base = datetime.fromisoformat(filas[0]["inicio"])
    assert [datetime.fromisoformat(f["inicio"]) - base - timedelta(weeks=i) for i, f in enumerate(filas)] == [timedelta(0)] + [timedelta(hours=2)] * 3
    # Historial: quién, cuándo, el motivo y qué cambió.
    h = admin.get(f"/api/agenda/{filas[2]['id']}").json()["historial"]
    assert len(h) == 1 and h[0]["motivo"] == "Cambió la sala" and h[0]["cambios"][0]["campo"] == "inicio"
    assert admin.get(f"/api/agenda/{filas[0]['id']}").json()["historial"] == []
    # Cancelar esta y las siguientes corta la serie; la primera sigue.
    c = admin.post(f"/api/agenda/{filas[2]['id']}/resultado", json={"estado": "cancelado", "resultado": "Se suspende el ciclo", "aplicar_a": "siguientes"})
    assert c.status_code == 200
    estados = [f["estado"] for f in admin.get("/api/agenda", params={"desde": desde, "hasta": hasta}).json()["eventos"]]
    assert estados == ["programado", "programado", "cancelado", "cancelado"]
    # El resultado de a una: no se da por «realizadas» todas juntas.
    assert admin.post(f"/api/agenda/{filas[0]['id']}/resultado", json={"estado": "realizado", "aplicar_a": "siguientes"}).status_code == 400


def test_repeticion_invalida(admin):
    assert _ev(admin, repetir={"frecuencia": "siempre", "cantidad": 3}).status_code == 422
    assert _ev(admin, repetir={"frecuencia": "semanal"}).status_code == 422
    assert _ev(admin, repetir={"frecuencia": "semanal", "cantidad": 500}).status_code == 422
    fin_de_mes = admin.post("/api/agenda", json={"tipo": "otro", "titulo": "Cierre", "inicio": "2030-01-31T12:00:00+00:00", "repetir": {"frecuencia": "mensual", "cantidad": 3}})
    assert fin_de_mes.status_code == 201
    dias = [f["inicio"][:10] for f in admin.get("/api/agenda", params={"desde": "2030-01-01T00:00:00+00:00", "hasta": "2030-06-01T00:00:00+00:00"}).json()["eventos"]]
    assert dias == ["2030-01-31", "2030-02-28", "2030-03-31"]


def test_filtro_por_persona_del_area(admin, cliente):
    crear_usuario("tec1", "Técnico", AREA_ADMISION)
    crear_usuario("tec2", "Técnico", AREA_ADMISION)
    a, b = cliente("tec1"), cliente("tec2")
    _ev(a, titulo="De uno")
    _ev(b, titulo="De otro")
    ids = {u["usuario"]: u["id"] for u in admin.get("/api/admin/usuarios").json()}
    desde, hasta = _cuando(-48), _cuando(24 * 10)
    todas = a.get("/api/agenda", params={"desde": desde, "hasta": hasta, "ambito": "area"}).json()["eventos"]
    assert {e["titulo"] for e in todas} == {"De uno", "De otro"}
    solo = a.get("/api/agenda", params={"desde": desde, "hasta": hasta, "ambito": "area", "usuario_id": ids["tec2"]}).json()["eventos"]
    assert [e["titulo"] for e in solo] == ["De otro"]


def test_superpuestos_vencimientos_y_ausencias(admin, cliente, db):
    from datetime import date

    from sqlalchemy import select

    from sulam.models import Agente, Area, Documento, SolicitudPersonal, TipoDocumental, Usuario

    crear_usuario("tec1", "Técnico", AREA_ADMISION)
    crear_usuario("tec2", "Técnico", AREA_ADMISION)
    crear_usuario("otro", "Técnico", OTRA)
    a, o = cliente("tec1"), cliente("otro")
    u2 = db.execute(select(Usuario).where(Usuario.usuario == "tec2")).scalar_one()
    admin_u = db.execute(select(Usuario).where(Usuario.usuario == "tec1")).scalar_one()
    ag = Agente(nombre="Colega Prueba", usuario_id=u2.id, area_id=u2.area_id, periodos=[{"desde": None, "hasta": None}])
    db.add(ag)
    db.flush()
    hoy = date.today()
    db.add(SolicitudPersonal(agente_id=ag.id, tipo="licencia", clase="enfermedad", desde=hoy, hasta=hoy + timedelta(days=2), estado="aprobada", pedido_por=u2.id, motivo="reservado"))
    db.add(SolicitudPersonal(agente_id=ag.id, tipo="comision", desde=hoy + timedelta(days=5), hasta=hoy + timedelta(days=5), lugar="Concepción", estado="registrada", pedido_por=u2.id))
    db.add(SolicitudPersonal(agente_id=ag.id, tipo="licencia", clase="anual", desde=hoy, hasta=hoy, estado="pendiente", pedido_por=u2.id))
    db.commit()
    # Un documento que vence en pocos días en un legajo del técnico.
    p = nueva_persona(admin)
    caso = nuevo_caso(admin, p["id"])
    admin.post(f"/api/casos/{caso['id']}/asignaciones", json={"usuario_id": str(admin_u.id)})
    tipo = db.execute(select(TipoDocumental).order_by(TipoDocumental.id)).scalars().first()
    tipo.tiene_vencimiento = True
    db.commit()
    r = admin.post("/api/documentos", data={"tipo_documental_id": str(tipo.id), "titulo": "Certificado", "persona_id": p["id"], "vence": (hoy + timedelta(days=3)).isoformat()},
                   files={"archivo": ("c.pdf", b"%PDF-1.4\n%%EOF", "application/pdf")})
    assert r.status_code in (200, 201), r.text
    rango = {"desde": f"{hoy}T00:00:00+00:00", "hasta": f"{hoy + timedelta(days=15)}T00:00:00+00:00"}
    # El técnico ve en su área la licencia (sin decir que es por enfermedad) y la comisión; la pendiente no figura.
    s = a.get("/api/agenda/superpuestos", params={**rango, "ambito": "area"}).json()
    assert [(x["agente"], x["tipo_texto"]) for x in s["ausencias"]] == [("Colega Prueba", "Ausente"), ("Colega Prueba", "Comisión de servicio")]
    assert s["ausencias"][1]["lugar"] == "Concepción"
    # Solo «lo mío»: no aparecen las ausencias de otros.
    assert a.get("/api/agenda/superpuestos", params={**rango, "ambito": "mio"}).json()["ausencias"] == []
    # Otra área no ve las ausencias de esta.
    assert o.get("/api/agenda/superpuestos", params={**rango, "ambito": "area"}).json()["ausencias"] == []
    # El administrador ve el vencimiento; quien no tiene el legajo, no.
    v = admin.get("/api/agenda/superpuestos", params={**rango, "ambito": "todos"}).json()["vencimientos"]
    assert len(v) == 1 and v[0]["tipo"] == tipo.nombre and v[0]["persona"]["id"] == p["id"]
    assert o.get("/api/agenda/superpuestos", params={**rango, "ambito": "area"}).json()["vencimientos"] == []


def test_salida_guardada_conserva_su_recorrido(admin):
    import uuid

    personas = [nueva_persona(admin, numero_documento=str(2000000 + i), apellidos=["GOMEZ", "LOPEZ", "BENITEZ"][i], nombres=["Rosa", "Hector", "Ines"][i]) for i in range(3)]
    casos = [nuevo_caso(admin, p["id"]) for p in personas]
    sid, hora = str(uuid.uuid4()), _cuando(30)
    # El recorrido se guarda con el orden elegido (el 3.º, el 1.º, el 2.º).
    ids = []
    for orden, i in enumerate((2, 0, 1), 1):
        r = _ev(admin, tipo="salida", titulo="Salida a domicilio", inicio=hora, persona_id=personas[i]["id"], caso_id=casos[i]["id"], salida_id=sid, salida_orden=orden)
        assert r.status_code == 201, r.text
        ids.append(r.json()["id"])
    s = admin.get(f"/api/agenda/{ids[1]}/salida").json()
    assert [p["caso_id"] for p in s["paradas"]] == [casos[2]["id"], casos[0]["id"], casos[1]["id"]]
    assert [p["actual"] for p in s["paradas"]] == [False, True, False]
    # Una actividad que no es salida no tiene recorrido, y el agrupamiento solo vale para salidas.
    otra = _ev(admin).json()["id"]
    assert admin.get(f"/api/agenda/{otra}/salida").status_code == 400
    assert _ev(admin, tipo="reunion", titulo="Reunión", salida_id=sid).status_code == 422
    # Las salidas guardadas antes de esta mejora (sin identificador) se reconocen por agendarse juntas.
    viejas = [_ev(admin, tipo="salida", titulo="Salida a domicilio", inicio=_cuando(60), persona_id=personas[i]["id"], caso_id=casos[i]["id"]).json()["id"] for i in (1, 0)]
    s = admin.get(f"/api/agenda/{viejas[0]}/salida").json()
    assert [p["caso_id"] for p in s["paradas"]] == [casos[1]["id"], casos[0]["id"]]


def test_orden_de_comision_de_una_salida(admin, cliente):
    import io
    import uuid

    from docx import Document

    crear_usuario("tec1", "Técnico", AREA_ADMISION)
    crear_usuario("tec2", "Técnico", AREA_ADMISION)
    u1 = admin.get("/api/agenda/participantes").json()
    ids = {x["nombre"]: x["id"] for x in u1}
    personas = [
        nueva_persona(admin, numero_documento=str(3000000 + i), domicilio=f"Calle Ficticia {100 + i}", localidad="Yerba Buena", departamento="Yerba Buena", apellidos=["SECRETOVEGA", "SECRETORIOS"][i], nombres=["Rosa", "Hector"][i])
        for i in range(2)
    ]
    sid, hora = str(uuid.uuid4()), "2026-11-03T12:00:00+00:00"  # 09:00 en Tucumán
    paradas = [
        _ev(admin, tipo="salida", titulo="Salida a domicilio", inicio=hora, persona_id=p["id"], salida_id=sid, salida_orden=n, participantes=[ids["Tec1"], ids["Tec2"]]).json()["id"]
        for n, p in enumerate(personas, 1)
    ]
    r = admin.get(f"/api/agenda/{paradas[0]}/comision.docx")
    assert r.status_code == 200 and "wordprocessingml" in r.headers["content-type"]
    texto = "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)
    assert "SEÑOR/A Tec1" in texto and "SEÑOR/A Tec2" in texto
    assert "Calle Ficticia 100, Yerba Buena; Calle Ficticia 101, Yerba Buena" in texto and "Departamento Yerba Buena" in texto
    assert "Visitas domiciliarias a personas adultas mayores (2)" in texto and "3 de noviembre de 2026 a Hs. 09:00" in texto
    assert "SECRETO" not in texto and "Rosa" not in texto  # el objeto no lleva nombres de las personas visitadas
    # Un recordatorio no lleva orden; quien no ve la actividad no la descarga.
    rec = _ev(admin, tipo="recordatorio", titulo="Llamar").json()["id"]
    assert admin.get(f"/api/agenda/{rec}/comision.docx").status_code == 400
    assert cliente("tec1").get(f"/api/agenda/{rec}/comision.docx").status_code == 403


def test_orden_de_comision_regreso_lo_decide_quien_descarga(admin):
    import io

    from docx import Document

    e = _ev(admin, tipo="salida", titulo="Salida", inicio="2026-11-03T12:00:00+00:00", fin="2026-11-03T15:00:00+00:00").json()["id"]

    def texto(q=""):
        r = admin.get(f"/api/agenda/{e}/comision.docx{q}")
        assert r.status_code == 200, r.text
        return "\n".join(p.text for p in Document(io.BytesIO(r.content)).paragraphs)

    assert "regresar de esta comisión el 3 de noviembre de 2026 a Hs. 13:30" in texto("?regreso=13:30")
    assert "regresar de esta comisión sin retorno" in texto("?regreso=sin_retorno")
    assert "sin retorno" not in texto() and "regresar de esta comisión………" in texto()  # en blanco, aunque la actividad tenga fin
    assert admin.get(f"/api/agenda/{e}/comision.docx?regreso=25:99").status_code == 422
