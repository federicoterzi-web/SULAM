"""Salidas de las estadísticas: Excel (una hoja por bloque + detalle) y PDF con membrete.

También arma el Excel de incidencias de asistencia, con las mismas columnas que
la app «Asistencia DAM».
"""
from __future__ import annotations

import html
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import estadisticas as E
from . import filtros as FL
from .models import Area

FUENTE = "Arial"
AZUL = "1F4E79"
COLUMNAS = {
    "cantidad": "Cantidad",
    "pct": "%",
    "activos": "Activos",
    "sin_actividad": "Sin movimiento",
    "asignados": "Asignados",
    "cerrados": "Cerrados",
    "evoluciones": "Evoluciones",
    "sin_responsable": "Sin técnico/a",
    "personas": "Personas",
    "jornadas": "Jornadas",
    "con_incidencias": "Con incidencias",
    "pendientes": "Pendientes",
    "horas": "Horas sin justificar",
    "pedidos": "Pedidos",
    "dias": "Días",
    "fecha": "Publicada",
    "para": "Para",
    "destinatarios": "Destinatarios",
    "confirmaron": "Confirmaron",
    "area": "Área",
    "ingresados": "Ingresados",
}
# Tablas en las que cada caso/registro cuenta una sola vez: llevan fila de total.
SUMABLES = {
    "ingresos_mes", "evoluciones_mes", "ingresos_tipo", "pases_estado", "pases_tipo", "carpetas_area", "informes_plantilla",
    "informes_area", "documentos_tipo", "licencias", "justificativos", "estado_jornadas", "asistencia_areas", "areas",
    "estado", "area", "tipo_ingreso", "derivado_por", "edad", "genero", "departamento", "cud", "tipo_vivienda",
    "estado_funcional", "tiene_autonomia", "estado_civil", "tiene_hijos", "sin_confirmar_area",
}


def nombre_archivo(seccion: str, q: E.Consulta, extension: str) -> str:
    return f"estadisticas-{seccion}-{q.desde.isoformat()}-a-{q.hasta.isoformat()}.{extension}"


def _f(d: date) -> str:
    return d.strftime("%d/%m/%Y")


def _area(db: Session, q: E.Consulta) -> str:
    if q.area_id is None:
        return "Toda la Dirección"
    a = db.get(Area, q.area_id)
    return a.nombre if a else "—"


def _encabezado(db: Session, q: E.Consulta, seccion: str) -> list[tuple[str, str]]:
    filas = [
        ("Período", f"{_f(q.desde)} al {_f(q.hasta)}"),
        ("Alcance", _area(db, q)),
    ]
    if seccion == "casos":
        textos = FL.describir(db, "casos", q.filtros)
        filas.append(("Filtros", "; ".join(textos) if textos else "Sin filtros"))
        filas.append(("Universo", "Casos con fecha de ingreso en el período"))
    filas.append(("Generado", datetime.now().astimezone().strftime("%d/%m/%Y %H:%M")))
    return filas


# ---------------------------------------------------------------- Excel


class _Hoja:
    def __init__(self, ws: Worksheet):
        self.ws = ws
        self.fila = 1
        self.anchos: dict[int, int] = {}

    def celda(self, col: int, valor, *, negrita=False, color=None, fondo=None, formato=None, ajustar=False):
        c = self.ws.cell(row=self.fila, column=col, value=valor)
        c.font = Font(name=FUENTE, size=10, bold=negrita, color=color)
        if fondo:
            c.fill = PatternFill("solid", fgColor=fondo)
        if formato:
            c.number_format = formato
        if ajustar:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        largo = len(str(valor)) if valor is not None and not str(valor).startswith("=") else 8
        self.anchos[col] = max(self.anchos.get(col, 8), min(largo + 2, 60))
        return c

    def titulo(self, texto: str):
        self.celda(1, texto, negrita=True).font = Font(name=FUENTE, size=13, bold=True, color=AZUL)
        self.fila += 1

    def datos(self, pares: list[tuple[str, str]]):
        for k, v in pares:
            self.celda(1, k, negrita=True)
            self.celda(2, v, ajustar=len(str(v)) > 60)
            self.fila += 1
        self.fila += 1

    def resumen(self, items: list[dict]):
        self.celda(1, "Indicador", negrita=True, color="FFFFFF", fondo=AZUL)
        self.celda(2, "Valor", negrita=True, color="FFFFFF", fondo=AZUL)
        self.celda(3, "Detalle", negrita=True, color="FFFFFF", fondo=AZUL)
        self.fila += 1
        for r in items:
            self.celda(1, r["etiqueta"])
            v = r.get("valor")
            self.celda(2, v if v is not None else "—", formato="0.0" if isinstance(v, float) else None)
            detalle = r.get("nota") or (f"{r['pct']} % del total" if r.get("pct") is not None else "")
            self.celda(3, detalle)
            self.fila += 1
        self.fila += 1

    def tabla(self, t: dict):
        self.celda(1, t["titulo"], negrita=True).font = Font(name=FUENTE, size=11, bold=True, color=AZUL)
        self.fila += 1
        if t.get("nota"):
            self.celda(1, t["nota"], color="666666")
            self.fila += 1
        base_ref = None
        if t.get("base") is not None:
            self.celda(1, "Casos del período (base de los %)")
            self.celda(2, t["base"], negrita=True)
            base_ref = f"$B${self.fila}"
            self.fila += 1
        filas = t["filas"]
        cols = list(t.get("columnas") or ["cantidad"])
        if any("area" in f for f in filas):
            cols = ["area", *cols]
        con_pct = base_ref is not None and cols == ["cantidad"]
        encabezado = ["", *[COLUMNAS.get(c, c) for c in cols]] + (["% sobre la base"] if con_pct else [])
        for i, h in enumerate(encabezado, start=1):
            c = self.celda(i, h or "Valor", negrita=True, color="FFFFFF", fondo=AZUL)
            c.alignment = Alignment(horizontal="left" if i == 1 else "right")
        self.fila += 1
        primera = self.fila
        if not filas:
            self.celda(1, "Sin datos en el período.", color="666666")
            self.fila += 1
        for f in filas:
            self.celda(1, f["etiqueta"], ajustar=len(str(f["etiqueta"])) > 60)
            for j, col in enumerate(cols, start=2):
                v = f.get(col)
                if col == "pct" and v is not None:
                    self.celda(j, v / 100, formato="0.0%")
                elif col == "fecha" and v:
                    self.celda(j, date.fromisoformat(v), formato="dd/mm/yyyy")
                else:
                    self.celda(j, v, formato="0.0" if isinstance(v, float) else None)
            if con_pct:
                self.celda(len(cols) + 2, f"=IF({base_ref}=0,\"\",B{self.fila}/{base_ref})", formato="0.0%")
            self.fila += 1
        if filas and t["clave"] in SUMABLES:
            ultima = self.fila - 1
            self.celda(1, "Total", negrita=True)
            for j, col in enumerate(cols, start=2):
                if col in ("area", "fecha", "para", "pct"):
                    continue
                letra = get_column_letter(j)
                self.celda(j, f"=SUM({letra}{primera}:{letra}{ultima})", negrita=True)
            borde = Border(top=Side(style="thin", color="999999"))
            for j in range(1, len(encabezado) + 1):
                self.ws.cell(row=self.fila, column=j).border = borde
            self.fila += 1
        self.fila += 1

    def cerrar(self):
        for col, ancho in self.anchos.items():
            self.ws.column_dimensions[get_column_letter(col)].width = max(ancho, 12 if col > 1 else 30)


def _libro() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    return wb


def _guardar(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_casos(db: Session, q: E.Consulta) -> bytes:
    d = E.casos(db, q)
    wb = _libro()
    h = _Hoja(wb.create_sheet("Resumen"))
    h.titulo("Estadísticas de casos — Dirección de Adultos Mayores")
    h.datos(_encabezado(db, q, "casos"))
    h.resumen(d["resumen"])
    h.tabla(E.tabla("ingresos_mes", "Ingresos y cierres por mes", [{"etiqueta": m["etiqueta"], "ingresados": m["ingresados"], "cerrados": m["cerrados"]} for m in d["mensual"]],
                    columnas=["ingresados", "cerrados"]))
    h.cerrar()
    for s in d["secciones"]:
        h = _Hoja(wb.create_sheet(s["titulo"][:31].replace("/", "-")))
        h.titulo(s["titulo"])
        for t in s["tablas"]:
            h.tabla(t)
        h.cerrar()
    # Detalle: una fila por caso, con autofiltro.
    ws = wb.create_sheet("Detalle de casos")
    filas = E.detalle_casos(db, q)
    columnas = list(filas[0].keys()) if filas else ["Caso"]
    for j, c in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=j, value=c)
        celda.font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
    for i, f in enumerate(filas, start=2):
        for j, c in enumerate(columnas, start=1):
            celda = ws.cell(row=i, column=j, value=f[c])
            celda.font = Font(name=FUENTE, size=10)
            if isinstance(f[c], date):
                celda.number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{max(len(filas) + 1, 1)}"
    for j, c in enumerate(columnas, start=1):
        largo = max([len(c)] + [len(str(f[c] or "")) for f in filas[:300]])
        ws.column_dimensions[get_column_letter(j)].width = min(max(largo + 2, 10), 50)
    return _guardar(wb)


BLOQUES_INTERNO = (
    ("carga", "Carga de trabajo y tiempos"),
    ("asistencia", "Asistencia del personal"),
    ("circuitos", "Circuitos administrativos"),
    ("comunicacion", "Comunicación interna"),
)


def excel_interno(db: Session, q: E.Consulta) -> bytes:
    d = E.interno(db, q)
    wb = _libro()
    for clave, titulo in BLOQUES_INTERNO:
        bloque = d.get(clave)
        if not bloque:
            continue
        h = _Hoja(wb.create_sheet(titulo[:31]))
        h.titulo(f"{titulo} — Dirección de Adultos Mayores")
        enc = _encabezado(db, q, "interno")
        if clave == "asistencia" and bloque.get("recortado"):
            enc.insert(1, ("Asistencia calculada", f"desde el {_f(date.fromisoformat(bloque['periodo']['desde']))} (máximo un año)"))
        h.datos(enc)
        h.resumen(bloque["resumen"])
        for t in bloque["tablas"]:
            h.tabla(t)
        h.cerrar()
    return _guardar(wb)


# ---------------------------------------------------------------- Excel de incidencias (asistencia)

ENCABEZADO_INCIDENCIAS = [
    "Trabajador", "Fecha", "Turno desde", "Turno hasta", "Entrada", "Salida", "Tardanza min", "Salida anticipada min",
    "Recuperados min", "Diferencia técnica min", "Alertas", "Justificativos", "Revisión", "Minutos propuestos", "Responsable", "Fundamento",
]


def excel_incidencias(filas: list[dict], incidencias: dict[str, str], estados: dict[str, str], referencias: dict[int, str], desde: date, hasta: date) -> bytes:
    """Mismo criterio que la app original: jornadas con datos que tuvieron alguna incidencia o una revisión manual."""
    wb = _libro()
    ws = wb.create_sheet("Incidencias")
    for j, h in enumerate(ENCABEZADO_INCIDENCIAS, start=1):
        c = ws.cell(row=1, column=j, value=h)
        c.font = Font(name=FUENTE, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=AZUL)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    i = 1
    for f in filas:
        alertas = f["incidencias"] + f["incidencias_justificadas"]
        revision = f.get("revision") or {}
        if f["estado"] == "sin_datos" or not (alertas or (revision and not revision.get("automatica"))):
            continue
        i += 1
        evaluada = f.get("evaluada")
        valores = [
            f["nombre"], date.fromisoformat(f["dia"]), f.get("turno_desde"), f.get("turno_hasta"),
            (f.get("entrada") or "")[:5] or None, (f.get("salida") or "")[:5] or None,
            *[(round(f[k]) if evaluada else None) for k in ("minutos_tarde", "minutos_salida_anticipada", "minutos_recuperados", "minutos_diferencia")],
            "; ".join(incidencias.get(x, x) for x in alertas),
            "; ".join(f"{j['codigo']} #{j['id']} ({j['estado']})" + (f" ref. {referencias[j['id']]}" if referencias.get(j["id"]) else "") for j in f["justificativos"]),
            estados.get(f["estado"], f["estado"]), revision.get("minutos") or 0, revision.get("usuario") or "", revision.get("fundamento") or "",
        ]
        for j, v in enumerate(valores, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.font = Font(name=FUENTE, size=10)
            if isinstance(v, date):
                c.number_format = "dd/mm/yyyy"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:P{i}"
    for j, ancho in enumerate([28, 11, 10, 10, 9, 9, 10, 12, 11, 12, 40, 34, 18, 11, 22, 40], start=1):
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[1].height = 30
    nota = wb.create_sheet("Criterios")
    for r, texto in enumerate([
        f"Período: {_f(desde)} al {_f(hasta)}.",
        "Se listan las jornadas con marcas que tuvieron alguna incidencia (justificada o no) o una revisión manual.",
        "Criterios fijos: se ignoran los segundos; tolerancia de 10 minutos si se recupera a la salida el mismo día; ventana de 30 minutos antes y después del turno.",
        "Los minutos propuestos son una propuesta para evaluación: SULAM no ejecuta descuentos ni liquidación.",
    ], start=1):
        nota.cell(row=r, column=1, value=texto).font = Font(name=FUENTE, size=10)
    nota.column_dimensions["A"].width = 120
    return _guardar(wb)


# ---------------------------------------------------------------- PDF


def pdf_estadisticas(db: Session, q: E.Consulta, seccion: str) -> bytes:
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.platypus import CondPageBreak, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from .config import get_settings
    from .informes import MEMBRETE, _fuente

    normal, negrita, cursiva, _ = _fuente()
    ancho, alto = A4
    util = ancho - 40 * mm
    azul = colors.HexColor("#1f4e79")
    barra = colors.HexColor("#2a6fb0")
    st_titulo = ParagraphStyle("t", fontName=negrita, fontSize=14, leading=18, textColor=azul, spaceAfter=4)
    st_sub = ParagraphStyle("s", fontName=normal, fontSize=9, leading=12, textColor=colors.HexColor("#444444"))
    st_sec = ParagraphStyle("sec", fontName=negrita, fontSize=12, leading=15, textColor=azul, spaceBefore=10, spaceAfter=4)
    st_tab = ParagraphStyle("tab", fontName=negrita, fontSize=10, leading=13, spaceBefore=8, spaceAfter=2)
    st_nota = ParagraphStyle("nota", fontName=cursiva, fontSize=8, leading=10, textColor=colors.HexColor("#555555"), spaceAfter=3)
    st_celda = ParagraphStyle("celda", fontName=normal, fontSize=8.5, leading=10.5)
    esc = lambda t: html.escape(str(t if t is not None else "—"), quote=False)  # noqa: E731

    def kpis(items: list[dict]):
        celdas = []
        for r in items:
            v = r.get("valor")
            v = "—" if v is None else (f"{v:.1f}".replace(".", ",") if isinstance(v, float) else f"{v}")
            if r.get("unidad"):
                v += f" {r['unidad']}"
            extra = r.get("nota") or (f"{str(r['pct']).replace('.', ',')} % del total" if r.get("pct") is not None else "")
            celdas.append(Paragraph(f"<font size=15 color='#1f4e79'><b>{esc(v)}</b></font><br/>{esc(r['etiqueta'])}"
                                    + (f"<br/><font size=7 color='#666666'>{esc(extra)}</font>" if extra else ""), ParagraphStyle("k", parent=st_celda, leading=13)))
        filas = [celdas[i:i + 3] for i in range(0, len(celdas), 3)]
        filas[-1] += [""] * (3 - len(filas[-1]))
        t = Table(filas, colWidths=[util / 3] * 3)
        t.setStyle(TableStyle([
            ("BOX", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")), ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cccccc")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        return t

    def barras(filas: list[dict], clave: str = "cantidad"):
        datos = [f for f in filas if isinstance(f.get(clave), (int, float))][:14]
        if not datos:
            return None
        maximo = max(f[clave] for f in datos) or 1
        alto_fila, ancho_eti = 13, 62 * mm
        d = Drawing(util, alto_fila * len(datos) + 4)
        largo = util - ancho_eti - 18 * mm
        for i, f in enumerate(datos):
            y = d.height - (i + 1) * alto_fila
            etiqueta = str(f["etiqueta"])
            etiqueta = etiqueta if len(etiqueta) <= 42 else etiqueta[:40] + "…"
            d.add(String(ancho_eti - 4, y + 3, etiqueta, fontName=normal, fontSize=7.5, textAnchor="end"))
            w = max(0.5, largo * f[clave] / maximo)
            d.add(Rect(ancho_eti, y + 1.5, w, alto_fila - 4, fillColor=colors.HexColor("#9aa7b4") if f.get("sin_dato") else barra, strokeColor=None))
            texto = f"{f[clave]}" + (f"  ({str(f['pct']).replace('.', ',')} %)" if f.get("pct") is not None else "")
            d.add(String(ancho_eti + w + 3, y + 3, texto, fontName=normal, fontSize=7.5))
        return d

    def tabla_pdf(t: dict) -> list:
        partes: list = [Paragraph(esc(t["titulo"]), st_tab)]
        if t.get("nota"):
            partes.append(Paragraph(esc(t["nota"]), st_nota))
        filas = t["filas"]
        if not filas:
            partes.append(Paragraph("Sin datos en el período.", st_nota))
            return [KeepTogether(partes)]
        cols = t.get("columnas")
        if not cols:
            g = barras(filas)
            if g:
                partes.append(g)
            return [KeepTogether(partes)]
        cols = (["area"] if any("area" in f for f in filas) else []) + list(cols)
        st_enc = ParagraphStyle("enc", parent=st_celda, fontName=negrita, fontSize=8, leading=9.5, textColor=colors.white)
        encabezado = [""] + [Paragraph(esc(COLUMNAS.get(c, c)), st_enc) for c in cols]

        def valor(f, c):
            v = f.get(c)
            if c == "fecha" and v:
                return _f(date.fromisoformat(v))
            if c == "area" and v:
                return v.replace("Departamento ", "").replace("División ", "")
            return f"{v}".replace(".", ",") if isinstance(v, float) else v

        cuerpo = [[Paragraph(esc(f["etiqueta"]), st_celda)] + [Paragraph(esc(valor(f, c)), st_celda) for c in cols] for f in filas[:60]]
        primera = util * (0.28 if len(cols) > 3 else 0.45)
        con_area = cols[0] == "area"
        ancho_area = util * 0.22 if con_area else 0
        resto = (util - primera - ancho_area) / (len(cols) - (1 if con_area else 0))
        anchos = [primera] + ([ancho_area] if con_area else []) + [resto] * (len(cols) - (1 if con_area else 0))
        tb = Table([encabezado] + cuerpo, colWidths=anchos, repeatRows=1)
        tb.setStyle(TableStyle([
            ("FONT", (0, 0), (-1, 0), negrita, 8), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("BACKGROUND", (0, 0), (-1, 0), azul),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f5f8")]), ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#dddddd")),
        ]))
        partes.append(tb)
        if len(filas) > 60:
            partes.append(Paragraph(f"Se muestran 60 de {len(filas)} filas; el detalle completo está en el Excel.", st_nota))
        return partes

    cuerpo: list = []
    titulo = "Estadísticas de casos" if seccion == "casos" else "Estadísticas de funcionamiento interno"
    cuerpo.append(Paragraph(titulo, st_titulo))
    for k, v in _encabezado(db, q, seccion):
        cuerpo.append(Paragraph(f"<b>{esc(k)}:</b> {esc(v)}", st_sub))
    cuerpo.append(Spacer(1, 4 * mm))
    if seccion == "casos":
        d = E.casos(db, q)
        cuerpo.append(kpis(d["resumen"]))
        cuerpo.append(Paragraph("Ingresos y cierres por mes", st_sec))
        mens = [[Paragraph("<b>Mes</b>", st_celda), Paragraph("<b>Ingresados</b>", st_celda), Paragraph("<b>Cerrados</b>", st_celda)]] + [
            [m["etiqueta"], m["ingresados"], m["cerrados"]] for m in d["mensual"]
        ]
        tm = Table(mens, colWidths=[40 * mm, 30 * mm, 30 * mm], repeatRows=1)
        tm.setStyle(TableStyle([("FONT", (0, 1), (-1, -1), normal, 8.5), ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#dddddd")),
                                ("ALIGN", (1, 0), (-1, -1), "RIGHT")]))
        cuerpo.append(tm)
        for s in d["secciones"]:
            cuerpo.append(CondPageBreak(60 * mm))
            cuerpo.append(Paragraph(esc(s["titulo"]), st_sec))
            for t in s["tablas"]:
                cuerpo += tabla_pdf(t)
    else:
        d = E.interno(db, q)
        for clave, titulo_b in BLOQUES_INTERNO:
            bloque = d.get(clave)
            if not bloque:
                continue
            cuerpo.append(CondPageBreak(80 * mm))
            cuerpo.append(Paragraph(titulo_b, st_sec))
            if clave == "asistencia" and bloque.get("recortado"):
                cuerpo.append(Paragraph(f"Asistencia calculada desde el {_f(date.fromisoformat(bloque['periodo']['desde']))} (máximo un año).", st_nota))
            cuerpo.append(kpis(bloque["resumen"]))
            for t in bloque["tablas"]:
                cuerpo += tabla_pdf(t)

    pie = get_settings().pie_informes

    class Lienzo(Canvas):
        def __init__(self, *a, **k):
            super().__init__(*a, **k)
            self._paginas = []

        def showPage(self):
            self._paginas.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            total = len(self._paginas)
            for estado in self._paginas:
                self.__dict__.update(estado)
                self.saveState()
                if MEMBRETE.is_file():
                    w = 62 * mm
                    self.drawImage(str(MEMBRETE), 20 * mm, alto - 10 * mm - w * 387 / 1396, w, w * 387 / 1396, mask="auto")
                self.setFont(normal, 7.5)
                self.setFillColor(colors.HexColor("#666666"))
                self.drawCentredString(ancho / 2, 11 * mm, f"{pie} · SULAM · Página {self._pageNumber} de {total}")
                self.restoreState()
                super().showPage()
            super().save()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=34 * mm, bottomMargin=20 * mm,
                            title=titulo, author="Dirección de Adultos Mayores", creator="SULAM")
    doc.build(cuerpo, canvasmaker=Lienzo)
    return buf.getvalue()


__all__ = ["excel_casos", "excel_interno", "excel_incidencias", "pdf_estadisticas", "nombre_archivo", "select"]
