"""Brief en Word (.docx) de un evento o de una campaña, para el área de Comunicación del Ministerio."""

from __future__ import annotations

import io
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt, RGBColor
from PIL import Image, ImageOps

from .routers.participacion import fecha_larga

AREA = "Dirección de Adultos Mayores · División Participación y Sensibilización Ciudadana"
AZUL = RGBColor(0x1F, 0x4E, 0x79)


def _documento() -> Document:
    d = Document()
    sec = d.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    st = d.styles["Normal"]
    st.font.name = "Arial"
    st.font.size = Pt(10.5)
    return d


def _titulo(d: Document, texto: str, sub: str | None = None) -> None:
    p = d.add_paragraph()
    r = p.add_run(AREA)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    h = d.add_paragraph()
    r = h.add_run(texto)
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = AZUL
    if sub:
        s = d.add_paragraph()
        s.add_run(sub).font.size = Pt(11)


def _fila(tabla, etiqueta: str, valor: str | None) -> None:
    c = tabla.add_row().cells
    c[0].width, c[1].width = Cm(4), Cm(12.6)
    r = c[0].paragraphs[0].add_run(etiqueta)
    r.bold = True
    c[1].paragraphs[0].add_run(valor or "—")


def _datos(d: Document, filas: list[tuple[str, str | None]]) -> None:
    t = d.add_table(rows=0, cols=2)
    t.style = "Light List Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for etiqueta, valor in filas:
        _fila(t, etiqueta, valor)


def _seccion(d: Document, titulo: str) -> None:
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    r = p.add_run(titulo)
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = AZUL


def _imagen_jpeg(datos: bytes, ancho_px: int = 1400) -> io.BytesIO:
    img = ImageOps.exif_transpose(Image.open(io.BytesIO(datos))).convert("RGB")
    if img.width > ancho_px:
        img = img.resize((ancho_px, round(img.height * ancho_px / img.width)))
    salida = io.BytesIO()
    img.save(salida, "JPEG", quality=85)
    salida.seek(0)
    return salida


def _fotos(d: Document, fotos: list[tuple[str, bytes]], ancho_cm: float) -> None:
    if not fotos:
        return
    t = d.add_table(rows=1, cols=len(fotos))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for celda, (nombre, datos) in zip(t.rows[0].cells, fotos):
        try:
            celda.paragraphs[0].add_run().add_picture(_imagen_jpeg(datos), width=Cm(ancho_cm))
        except Exception:  # una imagen dañada no impide armar el brief
            celda.paragraphs[0].add_run(f"(no se pudo incluir la imagen {nombre})")


def _cuando(e) -> str:
    ini = e.inicio.astimezone()
    texto = fecha_larga(ini.date()).capitalize() + f" de {ini.year}"
    if e.todo_el_dia:
        return texto + " (todo el día)"
    texto += f", {ini.strftime('%H:%M')} hs"
    if e.fin:
        texto += f" a {e.fin.astimezone().strftime('%H:%M')} hs"
    return texto


def _organiza(x) -> str:
    return f"{x.organiza_otra} (otra organización)" if x.organiza_otra else "Dirección de Adultos Mayores"


def _pie(d: Document, quien: str) -> None:
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run(f"Brief generado desde SULAM el {datetime.now().strftime('%d/%m/%Y %H:%M')} por {quien}.")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)


def _a_bytes(d: Document) -> bytes:
    salida = io.BytesIO()
    d.save(salida)
    return salida.getvalue()


def brief_de_evento(e, x, fotos: list[tuple[str, bytes]], quien: str) -> bytes:
    d = _documento()
    _titulo(d, e.titulo, f"{x.tipo.nombre} · {_cuando(e)}")
    _datos(
        d,
        [
            ("Tipo de evento", x.tipo.nombre),
            ("Fecha y hora", _cuando(e)),
            ("Lugar", e.lugar),
            ("Organiza", _organiza(x)),
            ("Campaña", x.campania.nombre if x.campania else None),
            ("Instituciones", ", ".join(x.instituciones or []) or None),
            ("Asistentes", str(x.asistentes) if x.asistentes is not None else None),
        ],
    )
    if x.resumen:
        _seccion(d, "Resumen")
        d.add_paragraph(x.resumen)
    if e.resultado:
        _seccion(d, "Cómo se desarrolló")
        d.add_paragraph(e.resultado)
    if fotos:
        _seccion(d, "Fotos")
        _fotos(d, fotos, 7.8 if len(fotos) > 1 else 12)
    _pie(d, quien)
    return _a_bytes(d)


def brief_de_campania(c, filas: list[tuple], quien: str) -> bytes:
    """``filas``: (actividad, datos del evento, fotos) de cada evento de la campaña."""
    d = _documento()
    periodo = None
    if c.desde or c.hasta:
        periodo = " al ".join(x.strftime("%d/%m/%Y") for x in (c.desde, c.hasta) if x)
    _titulo(d, f"Campaña: {c.nombre}", periodo)
    realizados = [(e, x, f) for e, x, f in filas if e.estado == "realizado"]
    proximos = [(e, x, f) for e, x, f in filas if e.estado == "programado"]
    _datos(
        d,
        [
            ("Estado", "Activa" if c.estado == "activa" else "Cerrada"),
            ("Eventos realizados", str(len(realizados))),
            ("Eventos por realizar", str(len(proximos))),
            ("Asistentes en total", str(sum(x.asistentes or 0 for _, x, _ in realizados))),
        ],
    )
    if c.descripcion:
        _seccion(d, "Descripción")
        d.add_paragraph(c.descripcion)
    if realizados:
        _seccion(d, "Eventos realizados")
        for e, x, fotos in realizados:
            p = d.add_paragraph()
            r = p.add_run(e.titulo)
            r.bold = True
            d.add_paragraph(f"{x.tipo.nombre} · {_cuando(e)}" + (f" · {e.lugar}" if e.lugar else ""))
            detalles = [f"Organiza: {_organiza(x)}"]
            if x.instituciones:
                detalles.append("Instituciones: " + ", ".join(x.instituciones))
            if x.asistentes is not None:
                detalles.append(f"Asistentes: {x.asistentes}")
            d.add_paragraph(" · ".join(detalles))
            if x.resumen:
                d.add_paragraph(x.resumen)
            _fotos(d, fotos, 6.5)
    if proximos:
        _seccion(d, "Próximos eventos")
        for e, x, _ in proximos:
            d.add_paragraph(f"{e.titulo} — {_cuando(e)}" + (f" · {e.lugar}" if e.lugar else "") + f" · Organiza: {_organiza(x)}", style="List Bullet")
    _pie(d, quien)
    return _a_bytes(d)
